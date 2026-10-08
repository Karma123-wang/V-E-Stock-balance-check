"""
Downloadable reports for the Medicine Stock Review: Excel, PDF and interactive HTML.

All three are built from the same content (build_content), so they always show the same numbers.
"""

from __future__ import annotations

import html
import io
import json
from datetime import date

import pandas as pd

import stock_core as sc

STATUS_COLOURS = {sc.EXPIRED: "F8CBAD", sc.EXP_3M: "FCE4D6", sc.EXP_6M: "FFF2CC"}
SHORT_NOTE = {
    sc.EXPIRY_CONFLICT: "Same batch with two expiry dates – correct in e-BMSIS",
    sc.EXACT_DUP: "Same batch, expiry and supply type – likely double entry",
    sc.SAME_BATCH_TYPES: "Same batch under Annual + Additional/Mobilized – confirm receipts",
    sc.MULTI_BATCH: "Different batches – normal, issue earliest expiry first (FEFO)",
}
RULE = ("Duplicate = the same medicine name AND the same dosage listed on more than one line. "
        "The same medicine with a different dosage (e.g. Amoxicillin 250mg capsule and 500mg capsule) "
        "is a separate item, not a duplicate.")


# ==================================================================================
# Shared content
# ==================================================================================
def build_content(res: dict, report_date: date, facility: str = "") -> dict:
    s = res["summary"]
    qty = bool(s["has_qty"])

    sections = [
        ("Items", [
            ("Stock lines in the report", s["lines"], "Every row in e-BMSIS"),
            ("Medicine names", s["medicine_names"], "Counting each medicine name once"),
            ("Total items (medicine + dosage)", s["medicines"],
             "Each dosage counted as its own item"),
            ("Medicines with more than one dosage", s["names_multi_dosage"],
             "Counted separately – NOT duplicates"),
        ]),
        ("Duplicates (same medicine + same dosage)", [
            ("Items listed more than once", s["medicines_repeated"], "See Duplicate Summary"),
            ("Lines belonging to these items", s["dup_lines"], ""),
            ("Extra lines (beyond one line per item)", s["extra_lines"], ""),
            *[(f"  {k}", v, SHORT_NOTE[k]) for k, v in s["dup_counts"].items()],
        ]),
        ("Expiry (lines)", [(k, v, "") for k, v in s["expiry_counts"].items()]),
    ]
    if qty:
        sections.append(("Stock level (items)", [(k[2:], v, "") for k, v in sorted(s["stock_counts"].items())]))

    notes = [RULE]
    if not qty:
        notes.append("The upload had no Balance Qty column, so stock levels (zero / low / overstock) "
                     "were not checked. Use the e-BMSIS Excel export to include them.")

    def pick(df, cols):
        cols = [(c, h) for c, h in cols if c in df.columns and not (c == "balance_qty" and not qty)]
        out = df[[c for c, _ in cols]].copy()
        out.columns = [h for _, h in cols]
        if "Expiry Date" in out:
            out["Expiry Date"] = pd.to_datetime(out["Expiry Date"]).dt.date
        return out.reset_index(drop=True)

    dup_sum = pick(res["dup_summary"], [
        ("medicine", "Medicine"), ("dosage", "Dosage"), ("lines", "Lines"), ("extra_lines", "Extra Lines"),
        ("batches", "Batches"), ("supply_types", "Supply Types"), ("finding", "Main Finding"),
        ("expired_lines", "Expired Lines"), ("rows", "Rows")])
    dup_detail = pick(res["duplicates"], [
        ("group_id", "Group"), ("medicine", "Medicine"), ("dosage", "Dosage"), ("row_no", "Row #"),
        ("supply_type", "Supply Type"), ("batch", "Batch No."), ("expiry", "Expiry Date"),
        ("expiry_status", "Expiry Status"), ("balance_qty", "Balance Qty"),
        ("duplicate_finding", "Finding"), ("action", "What to do")])
    multi = pick(res["name_dosage"], [
        ("medicine", "Medicine"), ("dosage", "Dosage"), ("lines", "Lines"),
        ("duplicated", "Duplicated?"), ("rows", "Rows")])
    expiry = pick(res["expiry"], [
        ("row_no", "Row #"), ("medicine", "Medicine"), ("dosage", "Dosage"), ("batch", "Batch No."),
        ("supply_type", "Supply Type"), ("expiry", "Expiry Date"), ("days_to_expiry", "Days Left"),
        ("expiry_status", "Expiry Status"), ("balance_qty", "Balance Qty"), ("action", "What to do")])
    all_items = pick(res["all"], [
        ("row_no", "Row #"), ("supply_type", "Supply Type"), ("medicine", "Medicine"), ("dosage", "Dosage"),
        ("batch", "Batch No."), ("expiry", "Expiry Date"), ("days_to_expiry", "Days Left"),
        ("expiry_status", "Expiry Status"), ("balance_qty", "Balance Qty"),
        ("duplicate_finding", "Duplicate Finding")])

    tables = [
        ("Duplicate Summary", "One row per medicine + dosage that is listed more than once.", dup_sum),
        ("Duplicate Detail", "Every line of the duplicated items, with what to do.", dup_detail),
        ("Same Name, Different Dosage", "Medicines stocked in more than one dosage. "
         "These are counted as separate items, not duplicates.", multi),
        ("Expiry Alerts", "Expired, or expiring within 6 months of the report date.", expiry),
    ]
    if qty and res["stock"] is not None:
        st = res["stock"].copy()
        st["stock_status"] = st["stock_status"].str[2:]
        cols = [("item_name", "Item (medicine + dosage)"), ("balance_qty", "Usable Balance Qty"),
                ("expired_qty", "Expired Qty"), ("monthly_use", "Monthly Use"),
                ("months_of_stock", "Months of Stock"), ("stock_status", "Stock Status"), ("rows", "Rows")]
        tables.append(("Stock Level", "Usable quantity added across batches of the same item.", pick(st, cols)))
    tables.append(("All Items", "Every line in the report.", all_items))

    title = "Medicine Stock Review" + (f" – {facility}" if facility else "")
    return {"title": title, "date": report_date.strftime("%d/%m/%Y"), "sections": sections,
            "notes": notes, "tables": tables}


def _cell(v):
    """Value as text for PDF / HTML."""
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NA or v is pd.NaT:
        return ""
    if isinstance(v, (date, pd.Timestamp)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ==================================================================================
# Excel
# ==================================================================================
def excel_report(res: dict, report_date: date, facility: str = "") -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    c = build_content(res, report_date, facility)
    buf = io.BytesIO()
    widths = {"Medicine": 30, "Dosage": 36, "Item (medicine + dosage)": 46, "What to do": 55,
              "Finding": 30, "Main Finding": 30, "Duplicate Finding": 30, "Batch No.": 24,
              "Supply Types": 18, "Rows": 14, "Expiry Status": 17, "Duplicated?": 24}
    head_fill = PatternFill("solid", fgColor="1F4E78")

    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        ws = xw.book.create_sheet("Summary")
        ws["A1"] = c["title"]
        ws["A1"].font = Font(name="Arial", bold=True, size=14)
        ws["A2"] = f"Report date: {c['date']}"
        ws["A2"].font = Font(name="Arial", size=10, color="595959")
        r = 4
        for name, rows in c["sections"]:
            ws.cell(r, 1, name).font = Font(name="Arial", bold=True, color="1F4E78", size=11)
            r += 1
            for label, value, note in rows:
                ws.cell(r, 1, label).font = Font(name="Arial", size=10)
                v = ws.cell(r, 2, int(value))
                v.font = Font(name="Arial", size=10, bold=True)
                v.alignment = Alignment(horizontal="center")
                ws.cell(r, 3, note).font = Font(name="Arial", size=9, color="595959")
                r += 1
            r += 1
        for n in c["notes"]:
            ws.cell(r, 1, n).font = Font(name="Arial", size=10, italic=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 42
            r += 1
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 10
        ws.column_dimensions["C"].width = 60

        for name, _, df in c["tables"]:
            df.to_excel(xw, sheet_name=name[:31], index=False)
            t = xw.sheets[name[:31]]
            for row in t.iter_rows():
                for cell in row:
                    cell.font = Font(name="Arial", size=10)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for i, cell in enumerate(t[1], 1):
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = head_fill
                t.column_dimensions[get_column_letter(i)].width = widths.get(cell.value, 12)
                if cell.value == "Expiry Date":
                    for rr in range(2, t.max_row + 1):
                        t.cell(rr, i).number_format = "DD/MM/YYYY"
                if cell.value == "Expiry Status":
                    for rr in range(2, t.max_row + 1):
                        col = STATUS_COLOURS.get(t.cell(rr, i).value)
                        if col:
                            t.cell(rr, i).fill = PatternFill("solid", fgColor=col)
            t.freeze_panes = "A2"
            if t.max_row > 1:
                t.auto_filter.ref = t.dimensions
        if "Sheet" in xw.book.sheetnames:
            del xw.book["Sheet"]
        xw.book.move_sheet("Summary", offset=-xw.book.index(xw.book["Summary"]))
        xw.book.active = 0
    return buf.getvalue()


# ==================================================================================
# PDF
# ==================================================================================
def _pdf_text(s: str) -> str:
    # Built-in PDF fonts have no "≤" or Greek "μ"; use characters they do have.
    return s.replace("≤", "<=").replace("μg", "mcg").replace("µg", "mcg").replace("μ", "mc").replace("µ", "mc").replace("–", "-")


def pdf_report(res: dict, report_date: date, facility: str = "", include_all_items: bool = True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    c = build_content(res, report_date, facility)
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=_pdf_text(c["title"]))
    ss = getSampleStyleSheet()
    H1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=17, alignment=0, spaceAfter=2)
    H2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12.5, textColor=colors.HexColor("#1F4E78"),
                        spaceBefore=6, spaceAfter=3)
    BODY = ParagraphStyle("b", parent=ss["Normal"], fontSize=9, leading=11.5)
    SMALL = ParagraphStyle("s", parent=ss["Normal"], fontSize=7.4, leading=9)
    SMALL_B = ParagraphStyle("sb", parent=SMALL, fontName="Helvetica-Bold", textColor=colors.white)
    GREY = ParagraphStyle("g", parent=BODY, textColor=colors.HexColor("#595959"))
    P = lambda t, st=SMALL: Paragraph(html.escape(_pdf_text(t)), st)
    blue = colors.HexColor("#1F4E78")
    usable_w = page[0] - 24 * mm

    story = [Paragraph(html.escape(_pdf_text(c["title"])), H1),
             Paragraph(f"Report date: {c['date']}", GREY), Spacer(1, 6)]

    # summary: sections side by side would be cramped, so stack them in one table
    rows, style = [], [("FONT", (0, 0), (-1, -1), "Helvetica", 9),
                       ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                       ("ALIGN", (1, 0), (1, -1), "CENTER"),
                       ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9D9D9"))]
    for name, items in c["sections"]:
        style += [("BACKGROUND", (0, len(rows)), (-1, len(rows)), colors.HexColor("#DDEBF7")),
                  ("FONT", (0, len(rows)), (-1, len(rows)), "Helvetica-Bold", 9.5),
                  ("SPAN", (0, len(rows)), (-1, len(rows)))]
        rows.append([_pdf_text(name), "", ""])
        for label, value, note in items:
            style.append(("FONT", (1, len(rows)), (1, len(rows)), "Helvetica-Bold", 9.5))
            rows.append([_pdf_text(label), str(value), P(note, GREY)])
    t = Table(rows, colWidths=[95 * mm, 22 * mm, usable_w - 117 * mm])
    t.setStyle(TableStyle(style))
    story += [t, Spacer(1, 8)]
    for n in c["notes"]:
        story.append(Paragraph("<i>" + html.escape(_pdf_text(n)) + "</i>", BODY))
        story.append(Spacer(1, 3))

    # relative column widths per header
    weight = {"Medicine": 2.3, "Dosage": 3.0, "Item (medicine + dosage)": 4.5, "What to do": 4.2,
              "Finding": 2.3, "Main Finding": 2.3, "Duplicate Finding": 2.3, "Batch No.": 2.0,
              "Supply Types": 1.5, "Supply Type": 1.1, "Rows": 1.3, "Expiry Status": 1.4,
              "Expiry Date": 1.1, "Duplicated?": 1.8, "Row #": 0.6, "Group": 0.6, "Lines": 0.6,
              "Extra Lines": 0.7, "Batches": 0.7, "Expired Lines": 0.8, "Days Left": 0.8}
    for name, desc, df in c["tables"]:
        if name == "All Items" and not include_all_items:
            continue
        story.append(PageBreak())
        story.append(Paragraph(html.escape(_pdf_text(name)) + f"  <font size=9 color='#595959'>({len(df)} rows)</font>", H2))
        story.append(Paragraph(html.escape(_pdf_text(desc)), GREY))
        story.append(Spacer(1, 4))
        if df.empty:
            story.append(Paragraph("Nothing to show.", BODY))
            continue
        heads = list(df.columns)
        w = [weight.get(h, 1.0) for h in heads]
        widths = [usable_w * x / sum(w) for x in w]
        data = [[P(h, SMALL_B) for h in heads]]
        tstyle = [("BACKGROUND", (0, 0), (-1, 0), blue),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
                  ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
        status_col = heads.index("Expiry Status") if "Expiry Status" in heads else None
        group_col = heads.index("Group") if "Group" in heads else None
        for i, row in enumerate(df.itertuples(index=False), 1):
            vals = [_cell(v) for v in row]
            data.append([P(v) for v in vals])
            if status_col is not None and row[status_col] in STATUS_COLOURS:
                tstyle.append(("BACKGROUND", (status_col, i), (status_col, i),
                               colors.HexColor("#" + STATUS_COLOURS[row[status_col]])))
            if group_col is not None and int(row[group_col]) % 2 == 0:
                tstyle.append(("BACKGROUND", (0, i), (status_col - 1 if status_col else -1, i),
                               colors.HexColor("#F2F2F2")))
        tbl = Table(data, colWidths=widths, repeatRows=1)
        tbl.setStyle(TableStyle(tstyle))
        story.append(tbl)

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#7F7F7F"))
        canvas.drawString(12 * mm, 7 * mm, _pdf_text(f"{c['title']}  ·  report date {c['date']}"))
        canvas.drawRightString(page[0] - 12 * mm, 7 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ==================================================================================
# Interactive HTML (one self-contained file: opens in any browser, works offline)
# ==================================================================================
def html_report(res: dict, report_date: date, facility: str = "") -> bytes:
    c = build_content(res, report_date, facility)
    s = res["summary"]
    data = {
        "title": c["title"], "date": c["date"], "notes": c["notes"],
        "sections": [[n, [[l.strip(), int(v), note] for l, v, note in rows]] for n, rows in c["sections"]],
        "tables": [{"name": n, "desc": d, "columns": list(df.columns),
                    "rows": [[_cell(v) for v in r] for r in df.itertuples(index=False)]}
                   for n, d, df in c["tables"]],
        "charts": {
            "Duplicate lines by finding": [[k, v] for k, v in s["dup_counts"].items()],
            "Lines by expiry status": [[k, v] for k, v in s["expiry_counts"].items()],
        },
    }
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    page = _HTML_TEMPLATE.replace("__TITLE__", html.escape(c["title"])).replace("__DATA__", payload)
    return page.encode("utf-8")


_HTML_TEMPLATE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__</title>
<style>
:root{--ink:#1d2433;--muted:#5b6475;--line:#e3e7ee;--bg:#f6f8fb;--card:#fff;--blue:#1f4e78;--blue2:#2e75b6;
--red:#f8cbad;--orange:#fce4d6;--yellow:#fff2cc;--chip:#eef3f9}
*{box-sizing:border-box}
body{margin:0;font:14px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif;color:var(--ink);background:var(--bg)}
header{background:var(--blue);color:#fff;padding:18px 24px}
header h1{margin:0;font-size:21px;font-weight:650}
header p{margin:4px 0 0;opacity:.85;font-size:13px}
main{max-width:1280px;margin:0 auto;padding:18px 16px 40px}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:16px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px}
.card .v{font-size:26px;font-weight:700;color:var(--blue)}
.card .l{font-size:12.5px;color:var(--muted)}
.card.warn .v{color:#c00000}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:12px;margin-bottom:16px}
.panel{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.panel h3{margin:0 0 10px;font-size:14px}
.bar{display:grid;grid-template-columns:minmax(150px,46%) 1fr 34px;gap:8px;align-items:center;margin:5px 0;font-size:12.5px}
.bar .track{background:var(--chip);border-radius:4px;height:14px;overflow:hidden}
.bar .fill{background:var(--blue2);height:100%}
.bar .n{text-align:right;font-weight:600}
table.sum{width:100%;border-collapse:collapse;font-size:13px}
table.sum td{padding:4px 6px;border-bottom:1px solid var(--line)}
table.sum td.v{text-align:center;font-weight:700;width:60px}
table.sum tr.sec td{background:#ddebf7;font-weight:650;color:var(--blue)}
table.sum td.note{color:var(--muted);font-size:12px}
.notes p{margin:6px 0;color:var(--muted);font-style:italic}
.tabs{display:flex;flex-wrap:wrap;gap:6px;margin:6px 0 10px}
.tabs button{border:1px solid var(--line);background:var(--card);border-radius:999px;padding:6px 12px;cursor:pointer;font:inherit;font-size:13px}
.tabs button.on{background:var(--blue);color:#fff;border-color:var(--blue)}
.tabs .cnt{opacity:.7;margin-left:4px}
.tools{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin-bottom:8px}
.tools input,.tools select{font:inherit;font-size:13px;padding:6px 9px;border:1px solid var(--line);border-radius:7px;background:#fff}
.tools input{min-width:240px;flex:1}
.desc{color:var(--muted);margin:0 0 8px;font-size:13px}
.shown{color:var(--muted);font-size:12.5px;margin-left:auto}
.wrap{overflow:auto;max-height:70vh;border:1px solid var(--line);border-radius:10px;background:#fff}
table.data{border-collapse:collapse;width:100%;font-size:12.5px}
table.data th{position:sticky;top:0;background:var(--blue);color:#fff;text-align:left;padding:7px 8px;cursor:pointer;white-space:nowrap;user-select:none}
table.data th .ar{opacity:.6;margin-left:3px}
table.data td{padding:6px 8px;border-bottom:1px solid var(--line);vertical-align:top}
table.data tr:hover td{background:#f3f7fc}
td.st-EXPIRED{background:var(--red)!important}
td.st-3{background:var(--orange)!important}
td.st-6{background:var(--yellow)!important}
td.flag{color:#c00000;font-weight:600}
@media print{.tools,.tabs{display:none}.wrap{max-height:none;overflow:visible}header{background:#fff;color:#000}}
</style></head>
<body>
<header><h1 id="ttl"></h1><p id="sub"></p></header>
<main>
  <div class="cards" id="cards"></div>
  <div class="grid2" id="charts"></div>
  <div class="grid2"><div class="panel"><h3>Summary</h3><table class="sum" id="sum"></table></div>
  <div class="panel notes"><h3>How to read this</h3><div id="notes"></div></div></div>
  <div class="tabs" id="tabs"></div>
  <p class="desc" id="desc"></p>
  <div class="tools">
    <input id="q" type="search" placeholder="Search medicine, dosage, batch…">
    <select id="f1"></select><select id="f2"></select>
    <span class="shown" id="shown"></span>
  </div>
  <div class="wrap"><table class="data" id="tbl"></table></div>
</main>
<script>
const D = __DATA__;
const $ = id => document.getElementById(id);
const esc = s => String(s).replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
$("ttl").textContent = D.title; $("sub").textContent = "Report date: " + D.date;

// cards
const flat = {}; D.sections.forEach(([n, rows]) => rows.forEach(([l, v]) => flat[l] = v));
const cards = [["Stock lines in the report","Stock lines"],["Medicine names","Medicine names"],
  ["Total items (medicine + dosage)","Total items (medicine + dosage)"],
  ["Medicines with more than one dosage","Medicines with several dosages"],
  ["Items listed more than once","Duplicated items",1],["EXPIRED","Expired lines",1],["Expires ≤3 months","Expiring ≤3 months",1]];
$("cards").innerHTML = cards.filter(c => c[0] in flat).map(([k,l,w]) =>
  `<div class="card${w && flat[k] ? " warn":""}"><div class="v">${flat[k]}</div><div class="l">${esc(l)}</div></div>`).join("");

// charts
$("charts").innerHTML = Object.entries(D.charts).map(([t, rows]) => {
  const mx = Math.max(1, ...rows.map(r => r[1]));
  return `<div class="panel"><h3>${esc(t)}</h3>` + rows.map(([l,v]) =>
    `<div class="bar"><span>${esc(l)}</span><div class="track"><div class="fill" style="width:${v/mx*100}%"></div></div><span class="n">${v}</span></div>`).join("") + "</div>";
}).join("");

// summary + notes
$("sum").innerHTML = D.sections.map(([n, rows]) => `<tr class="sec"><td colspan="3">${esc(n)}</td></tr>` +
  rows.map(([l,v,note]) => `<tr><td>${esc(l)}</td><td class="v">${v}</td><td class="note">${esc(note)}</td></tr>`).join("")).join("");
$("notes").innerHTML = D.notes.map(n => `<p>${esc(n)}</p>`).join("");

// tables
let cur = 0, sortCol = -1, sortDir = 1;
const FILTER_COLS = ["Main Finding","Finding","Duplicate Finding","Expiry Status","Supply Type","Duplicated?","Stock Status","Medicine"];
function tabs(){ $("tabs").innerHTML = D.tables.map((t,i) =>
  `<button class="${i===cur?"on":""}" data-i="${i}">${esc(t.name)}<span class="cnt">${t.rows.length}</span></button>`).join("");
  document.querySelectorAll("#tabs button").forEach(b => b.onclick = () => { cur = +b.dataset.i; sortCol = -1; setup(); }); }
function fillSelect(sel, t, avoid){
  const cols = t.columns.filter(c => FILTER_COLS.includes(c) && c !== avoid);
  const col = cols.find(c => c !== "Medicine") || cols[0];
  if (!col){ sel.style.display = "none"; sel.dataset.col = ""; return null; }
  const ci = t.columns.indexOf(col);
  const vals = [...new Set(t.rows.map(r => r[ci]).filter(v => v !== ""))].sort();
  if (vals.length < 2 || vals.length > 40){ sel.style.display = "none"; sel.dataset.col = ""; return null; }
  sel.style.display = ""; sel.dataset.col = ci;
  sel.innerHTML = `<option value="">All – ${esc(col)}</option>` + vals.map(v => `<option>${esc(v)}</option>`).join("");
  return col;
}
function setup(){
  tabs(); const t = D.tables[cur];
  $("desc").textContent = t.desc; $("q").value = "";
  const used = fillSelect($("f1"), t, null);
  const other = t.columns.filter(c => FILTER_COLS.includes(c) && c !== used && c !== "Medicine")[0];
  if (other){ const sel = $("f2"), ci = t.columns.indexOf(other);
    const vals = [...new Set(t.rows.map(r => r[ci]).filter(v => v !== ""))].sort();
    if (vals.length >= 2 && vals.length <= 40){ sel.style.display=""; sel.dataset.col = ci;
      sel.innerHTML = `<option value="">All – ${esc(other)}</option>` + vals.map(v => `<option>${esc(v)}</option>`).join("");
    } else { sel.style.display="none"; sel.dataset.col=""; }
  } else { $("f2").style.display = "none"; $("f2").dataset.col = ""; }
  render();
}
const num = v => { const s = String(v).replace(/,/g,""); return s !== "" && !isNaN(s) ? +s : null; };
const dkey = v => /^\d{2}\/\d{2}\/\d{4}$/.test(v) ? v.slice(6)+v.slice(3,5)+v.slice(0,2) : null;
function render(){
  const t = D.tables[cur], q = $("q").value.trim().toLowerCase();
  let rows = t.rows.filter(r => (!q || r.join(" ").toLowerCase().includes(q)) &&
    [$("f1"),$("f2")].every(s => !s.dataset.col || !s.value || r[+s.dataset.col] === s.value));
  if (sortCol >= 0) rows = rows.slice().sort((a,b) => {
    const x = a[sortCol], y = b[sortCol];
    const dx = dkey(x), dy = dkey(y); if (dx && dy) return dx < dy ? -sortDir : dx > dy ? sortDir : 0;
    const nx = num(x), ny = num(y); if (nx !== null && ny !== null) return (nx - ny) * sortDir;
    return String(x).localeCompare(String(y)) * sortDir; });
  const st = t.columns.indexOf("Expiry Status");
  const flagCols = ["Main Finding","Finding","Duplicate Finding"].map(c => t.columns.indexOf(c)).filter(i => i >= 0);
  $("tbl").innerHTML = "<thead><tr>" + t.columns.map((c,i) =>
      `<th data-i="${i}">${esc(c)}<span class="ar">${i===sortCol ? (sortDir>0?"▲":"▼") : "↕"}</span></th>`).join("") + "</tr></thead><tbody>" +
    rows.map(r => "<tr>" + r.map((v,i) => {
      let cls = "";
      if (i === st) cls = v === "EXPIRED" ? "st-EXPIRED" : v.includes("3 months") ? "st-3" : v.includes("6 months") ? "st-6" : "";
      if (flagCols.includes(i) && v && !v.startsWith("Different")) cls = "flag";
      return `<td class="${cls}">${esc(v)}</td>`; }).join("") + "</tr>").join("") + "</tbody>";
  $("shown").textContent = `Showing ${rows.length} of ${t.rows.length}`;
  document.querySelectorAll("#tbl th").forEach(th => th.onclick = () => {
    const i = +th.dataset.i; sortDir = i === sortCol ? -sortDir : 1; sortCol = i; render(); });
}
$("q").oninput = render; $("f1").onchange = render; $("f2").onchange = render;
setup();
</script>
</body></html>
"""
